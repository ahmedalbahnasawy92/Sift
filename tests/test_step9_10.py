"""
test_step9_10.py: offline tests for fetch_documents, read_table_cells (step 9)
and workflows (step 10). No LLM or API key.

Run:  python tests/test_step9_10.py
"""

import asyncio
import os
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))  # repo root

TMP = Path(tempfile.mkdtemp(prefix="step9_"))
os.environ["DOCS_ROOT"] = str(TMP / "docs")
os.environ["DOC_CACHE_DIR"] = str(TMP / "cache")
os.environ["SPOTLIGHT_SECRET_FILE"] = str(TMP / "secret")
os.environ["AUDIT_LOG"] = str(TMP / "audit.log")
os.environ["WORKFLOWS_DIR"] = str(TMP / "workflows")
os.environ["MAX_READ_CHARS"] = "20000"

U1 = TMP / "docs" / "u1"
U1.mkdir(parents=True)
(U1 / "a.txt").write_text("Contract A: fee AED 50,000. " + "x" * 100)
(U1 / "b.txt").write_text("Contract B: fee AED 65,000.")
(U1 / "big.txt").write_text("BIG " * 20000)

import openpyxl  # noqa: E402

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Budget"
ws.append(["Item", "Q1", "Q2", "Total"])
for i in range(1, 151):
    ws.append([f"Item {i}", i * 10, i * 20, None])
ws["D2"] = "=B2+C2"                      # formula: we must get the VALUE, not the text
wb.create_sheet("الموظفون").append(["الاسم", "الإدارة"])
wb["الموظفون"].append(["سالم راشد", "الموارد البشرية"])
wb.save(U1 / "budget.xlsx")
(U1 / "staff.csv").write_text("name,dept\nSalem,HR\nFatima,Finance\n")

# workflows: one shared, one user override, one invalid
for scope, name, desc in [("shared", "document-brief", "Shared brief."),
                          ("shared", "gazette-appointments", "List appointments."),
                          ("users/u1", "document-brief", "MY brief, overrides shared."),
                          ("shared", "Bad_Name", "Invalid because the name is uppercase.")]:
    d = TMP / "workflows" / scope / name
    d.mkdir(parents=True)
    (d / "SKILL.md").write_text(f"---\nname: {name}\ndescription: {desc}\n---\n\n1. Step one for {name}.\n")

from agno.run import RunContext  # noqa: E402

import doc_tools as d  # noqa: E402
import tool_guard as g  # noqa: E402
import workflows as w  # noqa: E402

CTX = RunContext(run_id="r1", session_id="s", user_id="u1")
CTX2 = RunContext(run_id="r1", session_id="s", user_id="u2")


def h(name):
    return next(x.doc_id for x in d._list_docs(CTX) if x.path.name == name)


def guarded(name, fn, args, ctx):
    async def call(**kw):
        return fn(**kw)
    return asyncio.run(g.tool_guard(name, call, {**args, "run_context": ctx}, ctx))


# ---------------------------------------------------------------- step 9
def test_1_fetch_several_documents():
    out = d.fetch_documents(f"{h('a.txt')}, {h('b.txt')}, doc-nope", CTX)
    assert "50,000" in out and "65,000" in out and "Not found: doc-nope" in out
    assert out.count("<<UNTRUSTED") >= 4  # names + bodies fenced
    print("1. fetch_documents reads several docs, reports missing ones    OK")


def test_2_fetch_shares_the_budget():
    out = d.fetch_documents(f"{h('big.txt')}, {h('a.txt')}", CTX)
    assert "TRUNCATED at 10,000" in out and "50,000" in out
    print("2. size budget shared between docs (big one truncated)        OK")


def test_3_fetch_and_read_dedupe_in_guard():
    ctx = RunContext(run_id="r-dedupe", session_id="s", user_id="u1")
    guarded("read_document", d.read_document, {"doc_id": h("a.txt")}, ctx)
    out = guarded("fetch_documents", d.fetch_documents, {"doc_ids": f"{h('a.txt')}, {h('b.txt')}"}, ctx)
    assert out.startswith("[Skipped, already read") and "65,000" in out and "50,000" not in out
    again = guarded("read_document", d.read_document, {"doc_id": h("b.txt")}, ctx)
    assert again == g.ALREADY_READ
    print("3. fetch skips docs already read; later re-read is blocked      OK")


def test_4_table_overview_and_header():
    out = d.read_table_cells(h("budget.xlsx"), CTX)
    assert "sheets ['Budget', 'الموظفون']" in out and "151 rows x 4 columns" in out
    assert "row | A | B | C | D" in out and "1 | Item | Q1 | Q2 | Total" in out
    print("4. overview: sheets, size, header row, column letters          OK")


def test_5_range_keeps_header_and_coordinates():
    out = d.read_table_cells(h("budget.xlsx"), CTX, cell_range="A50:C52")
    assert "1 | Item | Q1 | Q2   (header)" in out
    assert "50 | Item 49 | 490 | 980" in out and "52 | Item 51 | 510 | 1020" in out
    assert "Cite cells as 'Budget'!" in out
    print("5. range A50:C52 with header row and real row numbers          OK")


def test_6_paging_and_formula_values():
    out = d.read_table_cells(h("budget.xlsx"), CTX, max_rows=10)
    assert "more row(s) in this range: call again with cell_range starting at row 11" in out
    v = d.read_table_cells(h("budget.xlsx"), CTX, cell_range="D2:D2")
    assert "=B2+C2" not in v  # we show values (openpyxl data_only), never formula text
    print("6. long ranges are paged; formulas shown as values, not text   OK")


def test_7_arabic_sheet_csv_and_errors():
    ar = d.read_table_cells(h("budget.xlsx"), CTX, sheet="الموظفون")
    assert "سالم راشد" in ar and "الموارد البشرية" in ar
    csv = d.read_table_cells(h("staff.csv"), CTX, cell_range="A2:B3")
    assert "2 | Salem | HR" in csv and "3 | Fatima | Finance" in csv
    assert "is not a spreadsheet" in d.read_table_cells(h("a.txt"), CTX)
    assert "No sheet 'Nope'" in d.read_table_cells(h("budget.xlsx"), CTX, sheet="Nope")
    assert "Invalid cell_range" in d.read_table_cells(h("budget.xlsx"), CTX, cell_range="zz-top")
    print("7. Arabic sheet names, CSV, wrong type/sheet/range handled     OK")


# ---------------------------------------------------------------- step 10
def test_8_workflows_load_with_override_and_validation():
    wfs = w.load_workflows("u1")
    assert set(wfs) == {"document-brief", "gazette-appointments"}, set(wfs)  # Bad_Name rejected
    assert wfs["document-brief"][1] == "mine" and "MY brief" in wfs["document-brief"][0].description
    assert w.load_workflows("u2")["document-brief"][1] == "shared"
    print("8. shared + user workflows; user overrides; invalid skipped    OK")


def test_9_prompt_block_and_tools():
    block = w.workflows_block("u1")
    assert "gazette-appointments (shared): List appointments." in block
    assert "Step one" not in block  # only names + descriptions in the prompt
    out = w.read_workflow("gazette-appointments", CTX)
    assert "Step one for gazette-appointments" in out and "All earlier rules still apply" in out
    assert "No workflow named 'nope'" in w.read_workflow("nope", CTX)
    assert "document-brief (mine)" in w.list_workflows(CTX)
    print("9. prompt shows descriptions only; read_workflow loads steps   OK")


def test_10_user_folder_is_safe():
    assert w._user_folder("../u1").name == "_u1"
    assert w.load_workflows("../u1")["document-brief"][1] == "shared"  # can't read u1's
    print("10. user workflow folders can't be escaped with ../            OK")


if __name__ == "__main__":
    for name, fn in sorted(globals().items(), key=lambda kv: int(kv[0].split("_")[1]) if kv[0].startswith("test_") else 0):
        if name.startswith("test_"):
            fn()
    print("\nAll step-9/10 tests passed.")
