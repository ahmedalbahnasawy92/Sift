"""
doc_tools.py: STEP 4, Mike's document tools for the Agno agent.

  list_documents()                       -> handles + names, never content
  read_document(doc_id)                  -> full text with [Page N] markers (capped)
  find_in_document(doc_id, query, ...)   -> Ctrl+F: matches + context + page

Mike ideas kept:
  - The model only sees short handles (doc-1a2b3c), never file paths.
  - read_document returns the whole document; find_in_document for lookups.
  - Matching ignores case and extra whitespace, and returns the original
    wording, so the model can quote it (citations get checked in step 7).
  - Every result says which page the text is on.
Fixes over Mike:
  - read_document is CAPPED (Mike sends 1,000-page files whole and overflows).
  - Extracted text is cached for every file type (Mike re-extracts each turn).
  - PDFs are OCR'd with Nanonets OCR2 (like Mujeeb) when NANO_OCR_API_BASE_URL
    is set: every page by default, since Arabic text layers are often broken.
    Warm the cache once with:  python doc_tools.py --extract docs/<user_id>

Permissions: each user only sees docs/<user_id>/. The user_id comes from the
run (RunContext), never from the model, so the model can't ask for another
user's files.

Setup: pip install pymupdf python-docx openpyxl
"""

import hashlib
import os
import re
from dataclasses import dataclass
from pathlib import Path
from typing import List, Optional, Tuple

from agno.run import RunContext
from dotenv import load_dotenv

load_dotenv()  # OCR settings are read at import, before agno_agent loads .env

from spotlight import fence, fence_inline, nonce_for  # STEP 6

DOCS_ROOT = Path(os.getenv("DOCS_ROOT", "docs"))
CACHE_DIR = Path(os.getenv("DOC_CACHE_DIR", ".doc_cache"))
MAX_READ_CHARS = int(os.getenv("MAX_READ_CHARS", "100000"))  # ~25k tokens
SUPPORTED = {".pdf", ".docx", ".xlsx", ".txt", ".md", ".csv", ".json"}

# OCR for PDFs (Nanonets OCR2 behind an OpenAI-compatible endpoint, as in Mujeeb).
#   OCR_MODE=all      OCR every page. Arabic PDF text layers are often broken
#                     (reversed words, split ligatures), so this is the default.
#   OCR_MODE=missing  OCR only pages with no text layer (scanned pages).
#   OCR_MODE=off      text layer only.
# OCR is off when NANO_OCR_API_BASE_URL is not set. Results are cached per file.
OCR_URL = os.getenv("NANO_OCR_API_BASE_URL", "").rstrip("/")
OCR_MODEL = os.getenv("NANO_OCR_MODEL", "ocr")
OCR_API_KEY = os.getenv("NANO_OCR_API_KEY") or os.getenv("OPENAI_API_KEY", "none")
OCR_MAX_TOKENS = int(os.getenv("NANO_OCR_MAX_TOKENS", "4096"))
OCR_MODE = os.getenv("OCR_MODE", "all") if OCR_URL else "off"
OCR_DPI = int(os.getenv("OCR_DPI", "200"))
OCR_WORKERS = int(os.getenv("OCR_WORKERS", "4"))
OCR_MIN_CHARS = 50  # a page with less text than this counts as scanned
OCR_PROMPT = (
    "Extract the text from the above document as if you were reading it naturally. "
    "Return the tables in html format. Return the equations in LaTeX representation. "
    "If there is an image in the document and image caption is not present, add a small "
    "description of the image inside the <img></img> tag; otherwise, add the image caption "
    "inside <img></img>. Watermarks should be wrapped in brackets. Ex: "
    "<watermark>OFFICIAL COPY</watermark>. Page numbers should be wrapped in brackets. Ex: "
    "<page_number>14</page_number> or <page_number>9/22</page_number>. "
    "Prefer using ☐ and ☑ for check boxes."
)


# ---------------------------------------------------------------------------
# Document registry: stable handles per user
# ---------------------------------------------------------------------------
@dataclass
class Doc:
    doc_id: str
    path: Path
    rel: str   # path shown to the model, relative to the user's folder


def _user_dir(run_context: Optional[RunContext]) -> Path:
    user_id = (run_context.user_id if run_context else None) or "anonymous"
    safe = re.sub(r"[^A-Za-z0-9_.-]", "_", user_id).lstrip(".")  # no "../" tricks
    path = (DOCS_ROOT / (safe or "anonymous")).resolve()
    if path.parent != DOCS_ROOT.resolve():  # must be exactly one level below docs/
        raise PermissionError("Invalid user id")
    return path


def _list_docs(run_context: Optional[RunContext]) -> List[Doc]:
    try:
        base = _user_dir(run_context)
    except PermissionError:
        return []
    if not base.exists():
        return []
    docs = []
    for p in sorted(base.rglob("*")):
        if (p.is_file() and p.suffix.lower() in SUPPORTED and not p.name.startswith(".")
                and p.resolve().is_relative_to(base)):  # skip symlinks pointing outside
            rel = p.relative_to(base).as_posix()
            # Stable handle: same file -> same id across turns (Mike's slugs
            # are per turn; stable ids survive history and summaries).
            doc_id = "doc-" + hashlib.sha1(rel.encode()).hexdigest()[:6]
            docs.append(Doc(doc_id, p, rel))
    return docs


def _resolve(doc_id: str, run_context: Optional[RunContext]) -> Optional[Doc]:
    docs = _list_docs(run_context)
    for d in docs:  # accept the handle, or the exact file name as a fallback
        if doc_id in (d.doc_id, d.rel, d.path.name):
            return d
    return None


def _not_found(doc_id: str) -> str:
    return f"No document '{doc_id}'. Call list_documents to see the available doc_ids."


# ---------------------------------------------------------------------------
# Text extraction -> list of (page_label, text). Cached by file content hash.
# ---------------------------------------------------------------------------
def _ocr_image(img: bytes) -> str:
    """One page image (JPEG) -> markdown via the Nanonets OCR model (OpenAI-
    compatible vision endpoint). Same prompt as Mujeeb: tables as HTML."""
    import base64
    import httpx
    body = {
        "model": OCR_MODEL,
        "temperature": 0.0,
        "max_tokens": OCR_MAX_TOKENS,
        "messages": [{"role": "user", "content": [
            {"type": "image_url",
             "image_url": {"url": "data:image/jpeg;base64," + base64.b64encode(img).decode()}},
            {"type": "text", "text": OCR_PROMPT},
        ]}],
    }
    headers = {"Authorization": f"Bearer {OCR_API_KEY}"}
    for attempt in range(3):
        try:
            r = httpx.post(f"{OCR_URL}/chat/completions", json=body, headers=headers, timeout=180)
            r.raise_for_status()
            return r.json()["choices"][0]["message"]["content"] or ""
        except Exception:
            if attempt == 2:
                raise
    return ""


def _extract_pdf(path: Path) -> List[Tuple[str, str]]:
    """Text layer per page; OCR per OCR_MODE. A page whose OCR fails keeps its
    text layer, so one bad page never loses the whole document."""
    import pymupdf
    from concurrent.futures import ThreadPoolExecutor
    with pymupdf.open(path) as pdf:
        layer = [page.get_text() for page in pdf]
        if OCR_MODE == "all":
            todo = list(range(len(layer)))
        elif OCR_MODE == "missing":   # scanned pages only
            todo = [i for i, t in enumerate(layer) if len(t.strip()) < OCR_MIN_CHARS]
        else:
            todo = []

        # Render here (pymupdf is not thread-safe); only the HTTP calls run in parallel
        images = {i: pdf[i].get_pixmap(dpi=OCR_DPI).tobytes("jpeg", jpg_quality=85) for i in todo}

    def run(i):
        try:
            return i, _ocr_image(images[i])
        except Exception:
            return i, None

    texts = list(layer)
    if todo:
        with ThreadPoolExecutor(max_workers=OCR_WORKERS) as pool:
            for i, ocr in pool.map(run, todo):
                if ocr and ocr.strip():
                    texts[i] = ocr
    return [(f"Page {i}", t) for i, t in enumerate(texts, 1)]


def _extract(path: Path) -> List[Tuple[str, str]]:
    ext = path.suffix.lower()
    if ext == ".pdf":
        return _extract_pdf(path)
    if ext == ".docx":
        import docx
        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:  # tables as pipe rows
            for row in t.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        return [("Document", "\n".join(parts))]  # Word has no fixed pages
    if ext == ".xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        out = []
        for ws in wb.worksheets:
            rows = [" | ".join("" if v is None else str(v) for v in r)
                    for r in ws.iter_rows(values_only=True)]
            out.append((f"Sheet {ws.title}", "\n".join(rows)))
        return out
    return [("Document", path.read_text(encoding="utf-8", errors="replace"))]


def _pages(doc: Doc) -> List[Tuple[str, str]]:
    data = doc.path.read_bytes()
    key = hashlib.sha256(data).hexdigest()
    if doc.path.suffix.lower() == ".pdf" and OCR_MODE != "off":
        key += f"-ocr-{OCR_MODE}"  # OCR text is cached separately from the text layer
    cache = CACHE_DIR / f"{key}.txt"
    sep = "\n\x1e"  # record separator between pages in the cache file
    if cache.exists():
        raw = cache.read_text(encoding="utf-8").split(sep)
        return [tuple(r.split("\x1f", 1)) for r in raw if "\x1f" in r]
    pages = _extract(doc.path)
    CACHE_DIR.mkdir(exist_ok=True)
    cache.write_text(sep.join(f"{lbl}\x1f{txt}" for lbl, txt in pages), encoding="utf-8")
    return pages


def _render(pages: List[Tuple[str, str]]) -> Tuple[str, List[Tuple[int, str]]]:
    """Join pages with [Page N] markers. Returns text + (start_offset, label) list."""
    chunks, starts, pos = [], [], 0
    for label, text in pages:
        header = f"[{label}]\n"
        starts.append((pos, label))
        chunk = header + text.strip() + "\n\n"
        chunks.append(chunk)
        pos += len(chunk)
    return "".join(chunks), starts


def _label_at(starts: List[Tuple[int, str]], offset: int) -> str:
    return _page_span(starts, offset, 0)[2]


def _page_span(starts: List[Tuple[int, str]], offset: int, text_len: int) -> Tuple[int, int, str]:
    """(body_start, end, label) of the page containing offset."""
    label, begin, end = (starts[0][1] if starts else "Document"), 0, text_len
    for i, (start, lbl) in enumerate(starts):
        if start > offset:
            break
        label, begin = lbl, start + len(f"[{lbl}]\n")
        end = starts[i + 1][0] if i + 1 < len(starts) else text_len
    return begin, end, label


# ---------------------------------------------------------------------------
# Tools
# ---------------------------------------------------------------------------
def list_documents(run_context: RunContext) -> str:
    """List the documents the user has available, with their doc_id handles.
    Call this first when the user mentions a document, file, contract, report,
    or "my files", and you don't know the doc_id yet.
    """
    docs = _list_docs(run_context)
    if not docs:
        return "The user has no documents yet."
    n = nonce_for(run_context.session_id)
    lines = [f"{len(docs)} document(s):"]
    for d in docs:
        kb = max(1, d.path.stat().st_size // 1024)
        lines.append(f"- {d.doc_id}: {fence_inline(d.rel, n)} ({kb} KB)")  # names are untrusted
    return "\n".join(lines)


def documents_block(user_id: str, session_id: str, limit: int = 50) -> str:
    """STEP 6 + 8: the AVAILABLE DOCUMENTS list for the system prompt: handles,
    fenced names and a one-line summary per file (never the content).
    Newest first; long lists are cut."""
    ctx = RunContext(run_id="prompt", session_id=session_id, user_id=user_id)
    docs = sorted(_list_docs(ctx), key=lambda d: d.path.stat().st_mtime, reverse=True)
    if not docs:
        return "AVAILABLE DOCUMENTS: none. If the user asks about documents, tell them to add files."
    try:  # STEP 8: summaries come from the search index (if indexed)
        from search_index import get_summaries
        summaries = get_summaries(user_id)
    except Exception:
        summaries = {}
    n = nonce_for(session_id)
    lines = [f"AVAILABLE DOCUMENTS ({len(docs)}), newest first:"]
    for d in docs[:limit]:
        kb = max(1, d.path.stat().st_size // 1024)
        line = f"- {d.doc_id}: {fence_inline(d.rel, n)} ({kb} KB)"
        if d.rel in summaries:  # model-written from document text -> untrusted
            line += f" - {fence_inline(summaries[d.rel], n, limit=240)}"
        lines.append(line)
    if len(docs) > limit:
        lines.append(f"... {len(docs) - limit} more not shown: call list_documents to see all.")
    lines.append("Use these doc_ids directly with read_document / find_in_document / search_documents.")
    return "\n".join(lines)


def read_document(doc_id: str, run_context: RunContext) -> str:
    """Read the full text of one document, with [Page N] markers.
    Call this before answering questions about, summarising or quoting a
    document. Read each document at most ONCE per answer; after that use the
    text you already have, or find_in_document for a specific phrase.
    For very large documents the text is cut off: use find_in_document then.

    Args:
        doc_id: The handle from list_documents, e.g. "doc-1a2b3c".
    """
    doc = _resolve(doc_id, run_context)
    if not doc:
        return _not_found(doc_id)
    n = nonce_for(run_context.session_id)
    name = fence_inline(doc.rel, n)
    try:
        pages = _pages(doc)
    except Exception as e:
        return f"Could not read {name}: {e}"

    text, starts = _render(pages)
    if not any(t.strip() for _, t in pages):
        hint = ("OCR found no text either" if OCR_MODE != "off"
                else "OCR is off: set NANO_OCR_API_BASE_URL")
        return (f"{name} has no extractable text. It may be a scanned PDF "
                f"({hint}). Tell the user.")

    header = f"Document {doc.doc_id}: {name} ({len(pages)} section(s), {len(text):,} chars)\n\n"
    note = ""
    if len(text) > MAX_READ_CHARS:
        last = _label_at(starts, MAX_READ_CHARS)
        note = (f"\n\n[TRUNCATED at {MAX_READ_CHARS:,} of {len(text):,} chars, in {last}. "
                "Use find_in_document to look for specific terms in the rest.]")
        text = text[:MAX_READ_CHARS]
    # STEP 6: the body is untrusted data; our notes stay OUTSIDE the fence
    return header + fence(text, n, source=f"{doc.doc_id}") + note


# STEP 8: Arabic-tolerant matching. Each character is mapped 1:1 (or dropped,
# for diacritics/tatweel) and we keep the original index of every kept char,
# so a match in the normalised text points back to the exact original wording.
_DROP = set(chr(c) for c in list(range(0x0610, 0x061B)) + list(range(0x064B, 0x0660)) + [0x0670, 0x0640])
_FOLD = {"أ": "ا", "إ": "ا", "آ": "ا", "ٱ": "ا", "ى": "ي", "ة": "ه", "ؤ": "و", "ئ": "ي",
         **{chr(0x0660 + i): str(i) for i in range(10)}, **{chr(0x06F0 + i): str(i) for i in range(10)}}


def _norm_with_map(text: str) -> Tuple[str, List[int]]:
    out, idx = [], []
    for i, ch in enumerate(text):
        if ch in _DROP:
            continue
        out.append(_FOLD.get(ch, ch).lower())
        idx.append(i)
    return "".join(out), idx


def _find_in_one(doc: Doc, query: str, max_results: int, context_chars: int):
    """Returns (total_matches, [snippet lines]) for one document."""
    text, starts = _render(_pages(doc))
    norm, idx = _norm_with_map(text)
    qnorm, _ = _norm_with_map(query)
    words = [re.escape(w) for w in qnorm.split()]
    pattern = re.compile(r"\s+".join(words))
    matches = list(pattern.finditer(norm))
    lines = []
    for m in matches[:max_results]:
        o_start, o_end = idx[m.start()], idx[m.end() - 1] + 1   # back to original text
        page_start, page_end, label = _page_span(starts, o_start, len(text))
        a = max(page_start, o_start - context_chars)  # context stays on its page
        b = min(page_end, o_end + context_chars)
        lines.append(f"[{label}] ...{' '.join(text[a:b].split())}...")
    return len(matches), lines


def find_in_document(
    doc_id: str,
    query: str,
    run_context: RunContext,
    max_results: int = 20,
    context_chars: int = 80,
) -> str:
    """Find an exact word or phrase (like Ctrl+F) and return each match with
    surrounding text and its page. Matching ignores case, extra spaces,
    Arabic diacritics and letter variants (أ/إ/آ=ا, ى=ي, ة=ه, ٤٦=46).
    Use for targeted lookups (a clause, name, number, date) instead of reading
    a whole document. Pass doc_id "all" (or "") to search EVERY document.
    If nothing is found, try search_documents (it also matches by meaning).

    Args:
        doc_id: The handle, e.g. "doc-1a2b3c", or "all" to search all documents.
        query: The exact word or phrase to find, e.g. "termination".
        max_results: Maximum matches to return (default 20).
        context_chars: Characters of context on each side (default 80).
    """
    if not query.strip():
        return "Empty query. Give a word or phrase to search for."
    n = nonce_for(run_context.session_id)
    max_results = max(1, min(max_results, 50))
    context_chars = max(20, min(context_chars, 1500))

    all_docs = doc_id.strip().lower() in ("", "all", "*")
    if all_docs:
        docs = _list_docs(run_context)
        if not docs:
            return "The user has no documents yet."
    else:
        doc = _resolve(doc_id, run_context)
        if not doc:
            return _not_found(doc_id)
        docs = [doc]

    total, blocks, budget = 0, [], max_results
    for doc in docs:
        try:
            count, lines = _find_in_one(doc, query, budget, context_chars)
        except Exception as e:
            blocks.append(f"{doc.doc_id}: could not read ({e})")
            continue
        total += count
        if lines:
            label = f"{doc.doc_id} ({doc.rel})" if all_docs else ""
            blocks += [f"- {label + ' ' if label else ''}{line}" for line in lines]
            budget -= len(lines)
        if budget <= 0:
            break

    where = f"{len(docs)} documents" if all_docs else f"{docs[0].doc_id} ({fence_inline(docs[0].rel, n)})"
    if total == 0:
        return (f'No matches for "{query}" in {where}. Try a shorter phrase, another word '
                "form, or search_documents (matches by meaning).")
    header = f'{total} match(es) for "{query}" in {where}:'
    shown = max_results - max(budget, 0)
    more = f"\n({total - shown} more not shown; refine the query.)" if total > shown else ""
    # STEP 6: snippets are untrusted data; header/notes stay outside the fence
    return header + "\n" + fence("\n".join(blocks), n, source="find_in_document") + more


DOCUMENT_TOOLS = [list_documents, read_document, find_in_document]  # + search_documents (search_index.py)


if __name__ == "__main__":
    # Warm the text cache (runs OCR once) so the first question doesn't wait:
    #   python doc_tools.py --extract docs/demo-user
    import sys
    import time
    if len(sys.argv) == 3 and sys.argv[1] == "--extract":
        for p in sorted(Path(sys.argv[2]).iterdir()):
            if p.is_file() and p.suffix.lower() in SUPPORTED and not p.name.startswith("."):
                t0 = time.time()
                pages = _pages(Doc("", p, p.name))
                chars = sum(len(t) for _, t in pages)
                print(f"{p.name}: {len(pages)} section(s), {chars:,} chars, "
                      f"{time.time() - t0:.0f}s (OCR_MODE={OCR_MODE})")
    else:
        print("usage: python doc_tools.py --extract docs/<user_id>")
